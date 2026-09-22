"""Validated per-job documents, reply state and preparation tracker sheets."""
from __future__ import annotations
import hashlib
from pathlib import Path

from job_agent.config.settings import settings


def document_links(out=None, profile_hash=None):
    from job_agent.tracking.export import _read_json
    out = Path(out or settings.outputs_dir).resolve()
    profile_hash = profile_hash or _read_json(settings.profile_path, {}).get('profile_hash')
    result = {}
    for folder, column, pattern in [('interview_prep', 'Interview Prep', '{}.md'),
                                     ('cover_letters', 'Cover Letter', 'cover_{}.pdf')]:
        for item in _read_json(out/folder/'manifest.json', []):
            job_id = item.get('job_id', '')
            path = (out/folder/pattern.format(job_id)).resolve()
            if (path.parent == out/folder and path.is_file() and profile_hash
                    and item.get('profile_hash') == profile_hash
                    and hashlib.sha256(path.read_bytes()).hexdigest() == item.get('sha256')):
                result.setdefault(job_id, {})[column] = str(path)
    return result


def enrich_rows(rows, out, profile):
    from job_agent.tracking.export import _read_json
    from job_agent.tracking.inbox import latest_replies
    from job_agent.evaluation.gaps import missing_evidence
    links = document_links(out, profile.get('profile_hash'))
    manual = _read_json(out/'manual_applications.json', {})
    events = _read_json(out/'inbox_events.json', {})
    latest = latest_replies(events)
    outcomes = _read_json(out/'application_results.json', {})
    submitted = {i.get('job_id'): i for i in outcomes.get('successful', []) if i.get('status') == 'applied'}
    contacts = _read_json(out/'warm_contacts.json', {})
    for job_id, row in rows.items():
        for column in ('Interview Prep', 'Cover Letter'):
            row[column] = links.get(job_id, {}).get(column, '')
        row['Skills Found In Profile'] = '; '.join(missing_evidence(profile, row.get('Missing Skills', '').split(';')))
        row['Possible Contacts'] = '; '.join(f"{i['name']} — {i['title']} ({i['source_url']})" for i in contacts.get(job_id, []))
        if job_id in manual:
            row['Applied At'] = manual[job_id].get('at', '')
        elif job_id in submitted:
            row['Applied At'] = submitted[job_id].get('finished_at') or ''
        if job_id in latest:
            row['Status'] = latest[job_id]['status']
            row['Last Reply At'] = latest[job_id].get('received_at') or latest[job_id].get('observed_at')


def sync_tracker():
    """Add a review sheet without rewriting the original application history."""
    from openpyxl import load_workbook
    from job_agent.tracking.export import JobsCsvExporter, spreadsheet_text
    path = settings.outputs_dir/'applications_tracker.xlsx'
    if not path.exists():
        return
    rows = JobsCsvExporter().load()
    workbook = load_workbook(path)
    if 'Preparation & Outcomes' in workbook.sheetnames:
        del workbook['Preparation & Outcomes']
    sheet = workbook.create_sheet('Preparation & Outcomes')
    columns = ['Job ID', 'Company', 'Title', 'Status', 'Applied At', 'Last Reply At', 'Interview Prep', 'Cover Letter']
    sheet.append(columns)
    for row in rows.values():
        if row.get('Interview Prep') or row.get('Applied At') or row.get('Last Reply At'):
            sheet.append([spreadsheet_text(row.get(col, '')) for col in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    temporary = path.with_suffix('.tmp.xlsx')
    workbook.save(temporary)
    workbook.close()
    temporary.replace(path)
