"""Portable output must retain evidence and never attach an unrelated PDF."""
import csv
import hashlib
import io
import json
from zipfile import ZipFile

import pytest

from job_agent.config.schema import JobPosting
from job_agent.config.settings import settings
from job_agent.tailoring.regional import regional_policy, remote_eligibility
from job_agent.tracking.bundle import build_application_pack, validate_application_pack
from job_agent.tracking.export import JobsCsvExporter


def job(location="Remote - India", **values):
    return JobPosting(id="job1", title="Product Manager", company="Example", location=location,
                      job_url="https://example.com/jobs/1", source="jobicy", is_remote=True, **values)


@pytest.mark.parametrize("location,country,paper", [
    ("Bengaluru, India", "india", "a4"), ("US only", "usa", "us-letter"),
    ("London", "uk", "a4"), ("Canada", "canada", "us-letter"),
    ("Germany", "germany", "a4"), ("Remote worldwide", "international", "a4"),
    ("United States / Canada", "international", "a4"),
])
def test_regional_selection_never_guesses_one_of_multiple_markets(location, country, paper):
    policy = regional_policy(job(location))
    assert policy["country"] == country
    assert policy["paper"] == paper


def test_remote_eligibility_is_not_work_authorization():
    assert remote_eligibility(job("Remote - US only"), "India")[0] == "Location restricted"
    assert remote_eligibility(job("Remote - Europe"), "India")[0] == "Needs review"
    assert remote_eligibility(job("Worldwide"), "India")[0] == "Worldwide advertised"
    assert remote_eligibility(job("India / United States"), "India")[0] == "Location matches"


def prepare(monkeypatch, tmp_path):
    out = settings.outputs_dir
    folder = out / "tailored_resumes"
    folder.mkdir(parents=True)
    profile = tmp_path / "profile.json"
    profile.write_text(json.dumps({"profile_hash": "current", "work_authorization": {"current_country": "India"}}))
    monkeypatch.setattr(settings, "profile_path", profile)
    (out / "scraped_jobs.json").write_text(json.dumps([job().model_dump()]))
    contents = b"%PDF-1.7 test fixture"
    pdf = folder / "resume_job1.pdf"
    pdf.write_bytes(contents)
    record = {"job_id": "job1", "pdf_path": str(pdf), "profile_hash": "current",
              "pdf_sha256": hashlib.sha256(contents).hexdigest(), "validation_passed": True,
              "mode": "regional", "regional": {"paper": "a4"}, "target_country": "india"}
    (folder / "manifest.json").write_text(json.dumps([record]))
    return out, folder, record


def test_pack_roundtrip_has_relative_links_and_byte_exact_resume(monkeypatch, tmp_path):
    out, folder, _ = prepare(monkeypatch, tmp_path)
    (out / "quality_report.json").write_text(json.dumps({"grade_out_of_10": 8.5}), encoding="utf-8")
    pack = build_application_pack(out)
    assert validate_application_pack(pack) == {
        "jobs": 1, "resumes": 1, "supporting_documents": 0,
        "checked_documents": 1, "valid": True,
    }
    with ZipFile(pack) as archive:
        rows = list(csv.DictReader(io.StringIO(archive.read("jobs.csv").decode("utf-8-sig"))))
        assert rows[0]["Open Resume"] == "resumes/resume_job1.pdf"
        assert archive.read(rows[0]["Open Resume"]) == (folder / "resume_job1.pdf").read_bytes()
        assert "resumes/resume_job1.pdf" in archive.read("index.html").decode()
        assert not json.loads(archive.read("manifest.json"))["omitted_pdfs"]
        assert str(out) not in archive.read("jobs.csv").decode("utf-8-sig")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(archive.read('applications_tracker.xlsx')), read_only=True)
        assert workbook.sheetnames == ['Ready for review', 'Latest search', 'All saved jobs']
        workbook.close()
        assert 'Extract the entire ZIP' in archive.read('index.html').decode()
        assert 'applications_ready.csv' in archive.read('index.html').decode()
        assert json.loads(archive.read("quality_report.json"))["grade_out_of_10"] == 8.5


@pytest.mark.parametrize("change", ["tampered", "foreign_profile", "failed_check", "missing"])
def test_pack_refuses_untrusted_attachments(monkeypatch, tmp_path, change):
    out, folder, record = prepare(monkeypatch, tmp_path)
    if change == "tampered":
        (folder / "resume_job1.pdf").write_bytes(b"different PDF")
    elif change == "foreign_profile":
        record["profile_hash"] = "someone-else"
    elif change == "failed_check":
        record["validation_passed"] = False
    else:
        (folder / "resume_job1.pdf").unlink()
    (folder / "manifest.json").write_text(json.dumps([record]))
    with ZipFile(build_application_pack(out)) as archive:
        assert not any(name.endswith(".pdf") for name in archive.namelist())
        assert json.loads(archive.read("manifest.json"))["omitted_pdfs"]


def test_pack_validator_rejects_manifest_hash_mismatch(monkeypatch, tmp_path):
    out, _, _ = prepare(monkeypatch, tmp_path)
    pack = build_application_pack(out)
    broken = tmp_path / "broken.zip"
    with ZipFile(pack) as source, ZipFile(broken, "w") as target:
        for name in source.namelist():
            contents = source.read(name)
            if name == "resumes/resume_job1.pdf":
                contents = b"%PDF-1.7 tampered"
            target.writestr(name, contents)
    with pytest.raises(ValueError, match="hash mismatch"):
        validate_application_pack(broken)


def test_external_csv_text_is_not_executable_and_scores_sort_numerically(monkeypatch, tmp_path):
    out, _, _ = prepare(monkeypatch, tmp_path)
    exporter = JobsCsvExporter(outputs_dir=out, csv_path=out / "test.csv")
    exporter._write([{"Job ID": "low", "Title": "=HYPERLINK(\"bad\")", "Fit Score": "9.0"},
                     {"Job ID": "high", "Title": "+danger", "Fit Score": "10.0"}])
    rows = list(exporter.load().values())
    assert [row["Job ID"] for row in rows] == ["high", "low"]
    assert rows[0]["Title"].startswith("'+")
    assert rows[1]["Title"].startswith("'=")
