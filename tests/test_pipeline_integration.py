"""Offline six-stage integration and repeat-run regressions.

Only external sourcing is replaced; PDF parsing/compilation, scoring, artifact
handoffs, dry-run application handling, SQLite and Excel are exercised together.
"""

import json
from pathlib import Path

import openpyxl
import pypdf
import pytest

from job_agent.automation.pipeline import AutoApplyPipeline
from job_agent.config.schema import JobPosting, SearchParameters
from job_agent.evaluation.embedder import SemanticEmbedder
from job_agent.evaluation.pipeline import SemanticEvaluationPipeline
from job_agent.evaluation.reranker import LLMReranker
from job_agent.intake.parser import ResumeParser
from job_agent.sourcing.delta_store import DeltaStore
from job_agent.sourcing.scraper import OmnichannelScraper
from job_agent.tailoring.compiler import TypstResumeCompiler
from job_agent.tailoring.pipeline import ResumeTailoringPipeline
from job_agent.tailoring.rewriter import ResumeTailorer
from job_agent.tracking.cold_email import ColdEmailGenerator
from job_agent.tracking.pipeline import FallbackTrackingPipeline
from job_agent.tracking.tracker import MasterTracker
from job_agent.config.normalize import utc_now_iso
from tests.generate_test_resume import generate_sample_pdf


@pytest.fixture
def flow(tmp_path, monkeypatch):
    resume = generate_sample_pdf(tmp_path / "candidate.pdf")
    profile_path = tmp_path / "profile.json"
    profile = ResumeParser(provider="offline").parse(resume, output_path=profile_path)
    job = JobPosting(
        id="integration_cloud", title="Senior Cloud Infrastructure Engineer",
        company="Example Engineering", location="Remote", is_remote=True,
        job_url="https://example.com/jobs/cloud", source="greenhouse",
        # Company feeds must carry a posting date to pass the freshness window.
        date_posted=utc_now_iso(),
        description="Senior engineer with 4+ years experience in Python, Go, AWS, "
        "Kubernetes, Terraform, Redis, PostgreSQL, Docker and distributed systems.",
    )
    store = DeltaStore(tmp_path / "delta.db")
    scraper = OmnichannelScraper(
        search_params=SearchParameters(target_domains=["Cloud Engineer"], is_remote=True),
        delta_store=store,
    )
    monkeypatch.setattr(scraper, "scrape_job_boards", lambda: [job, job])
    jobs_path = tmp_path / "scraped_jobs.json"
    sourced = scraper.run_sourcing_pipeline(include_ats_direct=False, output_file=jobs_path)
    assert len(sourced) == 1
    embedder = SemanticEmbedder()
    embedder._use_fallback = True
    evaluation = SemanticEvaluationPipeline(
        embedder=embedder, reranker=LLMReranker(provider="none"), delta_store=store,
    )
    _, qualified = evaluation.run_evaluation(
        profile_path=profile_path, jobs_path=jobs_path, output_dir=tmp_path,
        tier1_threshold=0.0,
    )
    assert len(qualified) == 1
    compiler = TypstResumeCompiler(output_dir=tmp_path / "resumes")
    tailoring = ResumeTailoringPipeline(
        tailorer=ResumeTailorer(provider="none"), compiler=compiler, delta_store=store,
    )
    qualified_path = tmp_path / "qualified_jobs.json"
    records = tailoring.run_tailoring(profile_path=profile_path, qualified_jobs_path=qualified_path)
    assert len(records) == 1
    return dict(
        root=tmp_path, profile=profile, profile_path=profile_path, job=job, store=store,
        scraper=scraper, evaluation=evaluation, tailoring=tailoring,
        jobs_path=jobs_path, qualified_path=qualified_path, records=records,
        manifest_path=compiler.output_dir / "manifest.json",
    )


def apply(flow, agent=None):
    return AutoApplyPipeline(agent=agent, delta_store=flow["store"]).run_applications(
        profile_path=flow["profile_path"], qualified_jobs_path=flow["qualified_path"],
        manifest_path=flow["manifest_path"],
        output_path=flow["root"] / "application_results.json", dry_run=True,
    )


def tracking(flow):
    return FallbackTrackingPipeline(
        email_generator=ColdEmailGenerator(provider="none"), delta_store=flow["store"],
        tracker=MasterTracker(flow["root"] / "tracker.xlsx"),
    )


def track(flow, all_jobs=True):
    return tracking(flow).process_fallbacks(
        profile_path=flow["profile_path"], qualified_jobs_path=flow["qualified_path"],
        application_results_path=flow["root"] / "application_results.json",
        force_track_all=all_jobs,
    )


def test_six_stage_offline_flow_and_empty_rerun(flow):
    pdf = pypdf.PdfReader(flow["records"][0]["pdf_path"])
    text = " ".join(page.extract_text() for page in pdf.pages)
    assert flow["profile"].contact.full_name in text
    assert "$340k" in text
    successful, failed = apply(flow)
    assert not failed
    assert successful[0]["status"] == "dry_run"
    assert successful[0]["applied"] is False
    for _ in range(2):
        logged = track(flow)
        assert logged[0]["status"] == "DRY RUN"
    workbook = openpyxl.load_workbook(flow["root"] / "tracker.xlsx")
    assert workbook.active.max_row == 2
    assert workbook.active.cell(2, 5).value == "DRY RUN"
    workbook.close()
    assert flow["store"].status_counts() == {"tailored": 1}

    # A second sweep sees no new jobs. Every subsequently invoked phase must
    # replace its own old artifact, rather than quietly reuse the previous batch.
    assert flow["scraper"].run_sourcing_pipeline(
        include_ats_direct=False, output_file=flow["jobs_path"],
    ) == []
    assert flow["evaluation"].run_evaluation(
        profile_path=flow["profile_path"], jobs_path=flow["jobs_path"], output_dir=flow["root"],
    ) == ([], [])
    assert json.loads(flow["qualified_path"].read_text()) == []
    assert flow["tailoring"].run_tailoring(
        profile_path=flow["profile_path"], qualified_jobs_path=flow["qualified_path"],
    ) == []
    assert json.loads(flow["manifest_path"].read_text()) == []
    assert apply(flow) == ([], [])
    assert json.loads((flow["root"] / "application_results.json").read_text()) == {
        "successful": [], "failed": [],
    }


def test_tracking_preserves_submitted_status_and_excludes_it_from_outreach(flow):
    successful, _ = apply(flow)
    successful[0].update(status="applied", applied=True)
    (flow["root"] / "application_results.json").write_text(
        json.dumps({"successful": successful, "failed": []}), encoding="utf-8",
    )
    flow["store"].update_status(flow["job"].id, "applied")
    assert track(flow, all_jobs=False) == []
    assert track(flow)[0]["status"] == "APPLIED"
    assert flow["store"].status_counts() == {"applied": 1}


def test_application_results_survive_cleanup_failure(flow):
    from job_agent.automation.agent import AutoApplyAgent

    class BrokenCleanupAgent(AutoApplyAgent):
        def close(self):
            raise RuntimeError("cleanup failed")

    with pytest.raises(RuntimeError, match="cleanup failed"):
        apply(flow, BrokenCleanupAgent(delta_store=flow["store"]))
    results = json.loads((flow["root"] / "application_results.json").read_text())
    assert results["successful"][0]["job_id"] == flow["job"].id
