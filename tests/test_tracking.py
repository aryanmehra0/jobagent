"""Unit and integration tests for Phase 6: Fallback Tracking and Cold Email Outreach."""

from pathlib import Path
import json
import pytest
import openpyxl

from job_agent.config.schema import (
    CandidateProfile,
    ContactInfo,
    WorkAuthorization,
    Education,
    WorkExperience,
    SkillSet,
    JobPosting,
    EvaluatedJob,
    EvaluationScore,
)
from job_agent.tracking.cold_email import ColdEmailGenerator
from job_agent.tracking.tracker import MasterTracker
from job_agent.tracking.styler import get_priority_fill, FILL_HIGH_PRIORITY, FILL_MED_PRIORITY, FILL_LOW_PRIORITY
from job_agent.sourcing.delta_store import DeltaStore
from job_agent.tracking.pipeline import FallbackTrackingPipeline


@pytest.fixture
def candidate_profile() -> CandidateProfile:
    """Fixture providing candidate profile."""
    prof = CandidateProfile(
        contact=ContactInfo(
            full_name="Alex Rivera",
            email="alex.rivera@example.com",
            phone="+1-415-555-0142",
            location="San Francisco, CA",
        ),
        summary="Senior Distributed Systems and Cloud Engineer.",
        work_authorization=WorkAuthorization(
            citizenship=["United States"],
            current_country="United States",
            authorized_countries=["United States"],
            requires_sponsorship=False,
        ),
        experience=[
            WorkExperience(
                company="ScaleFlow Technologies",
                title="Lead Infrastructure Engineer",
                start_date="2022-01",
                description_bullets=[
                    "Architected multi-region Kubernetes clusters supporting over 200k RPS.",
                ],
            )
        ],
        skills=SkillSet(
            languages=["Python", "Go"],
            frameworks=["FastAPI"],
            cloud_devops=["Kubernetes", "AWS"],
        ),
        years_of_experience=4.0,
    )
    prof.seal_profile()
    return prof


@pytest.fixture
def target_job() -> JobPosting:
    """Fixture providing job posting."""
    return JobPosting(
        id="job_track_001",
        title="Staff Infrastructure Engineer",
        company="Stripe",
        location="Remote",
        job_url="https://stripe.com/jobs/infrastructure-1",
        description="Lead global payment infrastructure reliability.",
        is_remote=True,
        source="greenhouse",
    )


def test_cold_email_synthesis(candidate_profile, target_job):
    """Verify cold outreach email synthesis generates subject, verified facts, and CTA."""
    gen = ColdEmailGenerator()
    email_text = gen.generate_email(candidate_profile, target_job, fit_score=8.8)

    assert "Subject:" in email_text
    assert "Alex Rivera" in email_text
    assert "Stripe" in email_text
    assert "Kubernetes" in email_text or "200k RPS" in email_text


def test_openpyxl_master_tracker_formatting(target_job, tmp_path: Path):
    """Verify MasterTracker creates styled workbook, applies wrap_text, and priority fills."""
    excel_file = tmp_path / "test_tracker.xlsx"
    tracker = MasterTracker(excel_path=excel_file)

    sample_email = "Subject: Inquiry\n\nDear Team,\n\nI am writing regarding the open role."
    row_num = tracker.log_application(
        job=target_job,
        match_score=8.8,
        status="FALLBACK REQUIRED",
        cold_email=sample_email,
        failure_reason="Cloudflare Turnstile CAPTCHA detected",
    )

    assert excel_file.exists()
    assert row_num == 2

    # Load and inspect openpyxl cell styles
    wb = openpyxl.load_workbook(str(excel_file))
    ws = wb.active

    # Check headers
    assert ws.cell(row=1, column=2).value == "Job Title"
    assert ws.cell(row=1, column=8).value == "Personalized Cold Outreach Email"

    # Check data row
    assert ws.cell(row=2, column=2).value == "Staff Infrastructure Engineer"
    assert ws.cell(row=2, column=3).value == "Stripe"
    assert ws.cell(row=2, column=4).value == "8.8"

    # Check PatternFill color-coding on Score cell
    fill_color = str(ws.cell(row=2, column=4).fill.start_color.rgb)
    assert "D1FAE5" in fill_color  # Soft green priority fill for score >= 8.5

    # Check Alignment wrap_text on Cold Email cell (Column 8)
    email_alignment = ws.cell(row=2, column=8).alignment
    assert email_alignment.wrap_text is True


def test_fallback_tracking_pipeline_end_to_end(candidate_profile, target_job, tmp_path: Path):
    """Verify FallbackTrackingPipeline reads application results and generates master tracker."""
    excel_file = tmp_path / "applications_tracker.xlsx"
    tracker = MasterTracker(excel_path=excel_file)

    prof_file = tmp_path / "profile.json"
    prof_file.write_text(candidate_profile.model_dump_json(), encoding="utf-8")

    app_results_file = tmp_path / "application_results.json"
    app_data = {
        "successful": [],
        "failed": [
            {
                "id": target_job.id,
                "title": target_job.title,
                "company": target_job.company,
                "job_url": target_job.job_url,
                "error": "Timeout after 3 consecutive DOM attempts",
                "pdf_path": str(tmp_path / "resume.pdf"),
            }
        ],
    }
    app_results_file.write_text(json.dumps(app_data), encoding="utf-8")

    delta_db = tmp_path / "delta.db"
    store = DeltaStore(db_path=delta_db)
    store.mark_seen(target_job, status="tailored")

    # profile_path is passed explicitly; without it the pipeline would read the
    # user's real profile.json and the test would depend on repository state.
    pipeline = FallbackTrackingPipeline(tracker=tracker, delta_store=store)
    records = pipeline.process_fallbacks(
        application_results_path=app_results_file,
        qualified_jobs_path=tmp_path / "qualified_jobs.json",
        profile_path=prof_file,
    )

    assert len(records) == 1
    assert records[0]["company"] == "Stripe"
    assert records[0]["row"] == 2
    assert excel_file.exists()
